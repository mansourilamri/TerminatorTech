import tkinter as tk
from tkinter import ttk
from tkinter import filedialog, messagebox
import time
import os

try:
    from openpyxl import Workbook, load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Types de licence
# OS   = Open Source (code source libre)
# PROP = Proprietaire (licence commerciale)
# FREE = Gratuit (pas open source, mais gratuit)
# FREEM = Freemium (version gratuite limitee, payant pour toutes les fonctionnalites)

LICENCE_LABELS = {
    "OS":    "Open Source",
    "PROP":  "Proprietaire",
    "FREE":  "Gratuit",
    "FREEM": "Freemium",
}

LICENCE_COLORS = {
    "OS":    "#00ff88",
    "PROP":  "#ff6b6b",
    "FREE":  "#ffd700",
    "FREEM": "#ff9f43",
}

LABS = {
    "Laboratoire Arts": {
        "neon": "#00ff88", "icon": "ART",
        "sections": {
            "Tous les ordinateurs (Adobe CC)": [
                ("After Effect", "PROP"), ("Animate", "PROP"),
                ("Audition", "PROP"), ("Bridge", "FREE"),
                ("Dimension", "PROP"), ("DNG Converter", "FREE"),
                ("Dreamweaver", "PROP"), ("Illustrator", "PROP"),
                ("InDesign", "PROP"), ("Lightroom", "PROP"),
                ("Media Encoder", "PROP"), ("Photoshop", "PROP"),
                ("Prelude", "PROP"), ("Premiere", "PROP"),
            ],
            "PC Windows 11": [
                ("Antidote", "PROP"), ("Autodesk Autocad", "PROP"),
                ("Audacity", "OS"), ("Autodesk Maya", "PROP"),
                ("D5 (Rhino)", "PROP"), ("Ms Office", "PROP"),
                ("OpenShot", "OS"), ("Reaper (x64)", "PROP"),
                ("Revit", "PROP"), ("Rhino", "PROP"),
                ("Shotcut", "OS"), ("Sketchup", "FREEM"),
                ("Solidworks", "PROP"), ("Ultimaker Cura", "OS"),
            ],
            "MAC OS": [
                ("Antidote", "PROP"), ("Audacity", "OS"),
                ("DaVinci Resolve", "FREEM"), ("EndNote", "PROP"),
                ("GarageBand", "FREE"), ("Gimp", "OS"),
                ("Glyphs", "PROP"), ("Krita", "OS"),
                ("Ms Office", "PROP"), ("ShotCut", "OS"),
                ("Xmind 8", "FREEM"), ("Zotero", "OS"),
            ],
        },
    },
    "Laboratoire Communication": {
        "neon": "#00cfff", "icon": "COM",
        "sections": {
            "MAC OS X 13 (Ventura)": [
                ("Ableton Live 12", "PROP"), ("Adobe Acrobat DC", "PROP"),
                ("Adobe After Effects", "PROP"), ("Adobe Audition", "PROP"),
                ("Adobe Animate", "PROP"), ("Adobe Aero", "FREE"),
                ("Adobe Bridge", "FREE"), ("Adobe Character Animator", "PROP"),
                ("Adobe Dimension", "PROP"), ("Adobe Illustrator", "PROP"),
                ("Adobe InDesign", "PROP"), ("Adobe Incopy", "PROP"),
                ("Adobe Lightroom Classic", "PROP"),
                ("Adobe Media Encoder", "PROP"),
                ("Adobe Photoshop", "PROP"), ("Adobe Premiere Pro", "PROP"),
                ("Adobe Premiere Rush", "PROP"), ("Adobe XD", "PROP"),
                ("Antidote", "PROP"), ("Apercu", "FREE"),
                ("Arduino IDE", "OS"), ("Audacity", "OS"),
                ("Automator", "FREE"), ("Avid Media Composer", "PROP"),
                ("Balados", "FREE"), ("BBEdit", "PROP"),
                ("Blackmagic Proxy Generator Lite", "FREE"),
                ("Blackmagic RAW", "FREE"), ("Blender", "OS"),
                ("DaVinci Resolve", "FREEM"), ("Dictaphone", "FREE"),
                ("FaceTime", "FREE"), ("Firefox", "OS"),
                ("Freeform", "FREE"), ("Gimp", "OS"),
                ("GitHub Desktop", "OS"), ("GrandMA3", "PROP"),
                ("iMovie", "FREE"), ("Keynote", "FREE"),
                ("Mail", "FREE"), ("Max", "PROP"),
                ("Maxon Cinema 4D", "PROP"), ("Melodyne 5", "PROP"),
                ("Messages", "FREE"), ("Microsoft Excel", "PROP"),
                ("Microsoft PowerPoint", "PROP"),
                ("Microsoft Teams", "FREEM"), ("Microsoft Word", "PROP"),
                ("Microsoft Outlook", "PROP"),
                ("Microsoft One Note", "FREE"), ("Millumin5", "PROP"),
                ("Multidictionnaire", "PROP"), ("Musique", "FREE"),
                ("Notes", "FREE"), ("Numbers", "FREE"),
                ("OBS", "OS"), ("OneDrive", "FREEM"),
                ("Pages", "FREE"), ("Photo Booth", "FREE"),
                ("Photos", "FREE"), ("Plans", "FREE"),
                ("Pro Tools", "PROP"), ("QLab", "FREEM"),
                ("QuickTime Player", "FREE"), ("REAPER", "PROP"),
                ("Safari", "FREE"), ("Shotcut", "OS"),
                ("SketchUp", "FREEM"), ("Storyboarder", "OS"),
                ("TextEdit", "FREE"), ("TouchDesigner", "FREEM"),
                ("Transfert d'images", "FREE"), ("TwistedWave", "PROP"),
                ("VLC", "OS"), ("VueScan", "PROP"),
                ("Xcode", "FREE"),
                ("Youlean Loudness Meter", "FREEM"),
                ("zoom.us", "FREEM"), ("Zotero", "OS"),
            ],
        },
    },
    "Laboratoire Sciences": {
        "neon": "#bf5fff", "icon": "SCI",
        "sections": {
            "Bureautique": [
                ("Adobe Reader", "FREE"), ("Antidote", "PROP"),
                ("EndNote x20.4.1", "PROP"),
                ("Equitrac express client", "PROP"), ("Gimp", "OS"),
                ("LibreOffice", "OS"), ("MiKteX", "OS"),
                ("Microsoft Office 2021", "PROP"),
                ("Open office", "OS"), ("Paint.NET", "FREE"),
                ("NetSupportSchool", "PROP"), ("Texworks", "OS"),
                ("Zotero", "OS"),
            ],
            "Chimie": [
                ("BSL Analysis", "PROP"), ("Cn3D", "FREE"),
                ("Cytoscape", "OS"), ("ChemDraw Prime 22", "PROP"),
                ("Gaussian", "PROP"), ("GaussView", "PROP"),
                ("GenoPro", "PROP"), ("Molmol", "FREE"),
                ("Orbital viewer", "FREE"), ("PyMOL", "OS"),
                ("SimUtext (Simbio)", "PROP"),
            ],
            "Informatique": [
                ("Anaconda", "FREEM"), ("Android studio", "FREE"),
                ("Aptana Studio 3", "OS"),
                ("Apache NetBeans IDE", "OS"), ("BlueJ", "OS"),
                ("BlueGriffon", "OS"),
                ("Eclipse c, PHP, Java", "OS"), ("Emacs", "OS"),
                ("Edupython", "OS"), ("FileZilla", "OS"),
                ("Git", "OS"), ("Intelli IDEA Community", "OS"),
                ("IIs Express", "FREE"), ("Java", "FREE"),
                ("JEdit", "OS"), ("Jenkins", "OS"),
                ("Microsoft Visual Studio code", "FREE"),
                ("Mobaxterm", "FREEM"), ("Node.js", "OS"),
                ("Notepad++", "OS"),
                ("Oracle JDeveloper Studio", "FREE"),
                ("Oracle-OraClient10g home1", "PROP"),
                ("Pep", "OS"), ("Prolog-SWI", "OS"),
                ("PuTTY", "OS"), ("Python", "OS"),
                ("Pyscripter", "OS"), ("SQL Plus", "PROP"),
                ("SQL Developer", "FREE"), ("Spyder", "OS"),
                ("VirtualBox", "OS"), ("VNC Viewer", "FREEM"),
                ("Vi (editeur de texte)", "OS"), ("Win SCP", "OS"),
                ("X2go", "OS"), ("XMLmind", "FREEM"),
                ("7-Zip File Manager", "OS"),
            ],
            "Unix": [
                ("Alex (gem)", "OS"), ("AMC (Java seulement)", "OS"),
                ("Android-Studio", "FREE"), ("Ant", "OS"),
                ("Aptana Studio", "OS"), ("Arduino", "OS"),
                ("Atom", "OS"), ("Autotools", "OS"),
                ("Blender", "OS"), ("BlueGriffon", "OS"),
                ("BlueJ", "OS"), ("Build-essential", "OS"),
                ("Cairo", "OS"), ("Ccache", "OS"),
                ("ChimeraX", "FREEM"), ("Clang", "OS"),
                ("Clang-dev", "OS"), ("Clang-Tools", "OS"),
                ("Clang-Format", "OS"), ("Clang-Tidy", "OS"),
                ("Clion", "PROP"), ("Cmake", "OS"),
                ("CodeBlocks", "OS"), ("Cordova", "OS"),
                ("Cunit", "OS"), ("Curl", "OS"),
                ("Data display Debugger", "OS"), ("Devhelp", "OS"),
                ("Dia", "OS"), ("Docker", "FREEM"),
                ("Doctest", "OS"), ("Docutils", "OS"),
                ("Doxygen", "OS"), ("Eclipse", "OS"),
                ("Emacs", "OS"), ("FileZilla", "OS"),
                ("G++", "OS"), ("Gcc", "OS"),
                ("Gdb", "OS"), ("Gdebi", "OS"),
                ("Geany", "OS"), ("Gimp", "OS"),
                ("Git", "OS"),
            ],
            "Mathematique": [
                ("GeoGebra 6", "OS"), ("Graphic Calculus 2.0", "FREE"),
                ("Graphing Calculator", "PROP"),
                ("IBM SPSS Statistics", "PROP"), ("Inkscape", "OS"),
                ("JMP Pro", "PROP"), ("Krita", "OS"),
                ("MATLAB", "PROP"), ("R x64", "OS"),
                ("R i386", "OS"), ("Rstudio", "OS"),
                ("SAS 9.4 (francais)", "PROP"),
                ("Scilab (64 bit)", "OS"), ("Tinkerplot", "PROP"),
            ],
            "Autres Departements": [
                ("ArcGis", "PROP"), ("Gpower", "FREE"),
                ("Hot2000", "FREE"), ("HotPotatoes", "FREE"),
                ("ImageJ", "OS"), ("IoGas", "PROP"),
                ("Justinmind", "FREEM"), ("Kinovea", "OS"),
                ("Lego Mindstorms Edu EV3", "PROP"),
                ("Leapfrog", "PROP"), ("OpenShot video", "OS"),
                ("PHREEQC", "FREE"), ("ProteoWizard", "OS"),
                ("QGIS Desktop", "OS"), ("RETScreen 4", "FREE"),
                ("RETScreen Plus", "PROP"),
                ("Scratch-desktop", "OS"), ("Shotcut", "OS"),
                ("Ucinet", "PROP"), ("VLC Media Player", "OS"),
                ("Vortex10", "PROP"),
            ],
        },
    },
    "Sciences de l'education": {
        "neon": "#ff6b6b", "icon": "EDU",
        "sections": {
            "PC Windows 11": [
                ("7-Zip", "OS"), ("ActiveInspire", "PROP"),
                ("Adobe Digital Edition", "FREE"), ("Antidote", "PROP"),
                ("Algodoo", "FREEM"), ("Audacity", "OS"),
                ("Balabolka", "FREE"), ("Clavicom NG+", "FREE"),
                ("Dia", "OS"), ("EndNote", "PROP"),
                ("Foxit reader", "FREEM"),
                ("Lego Ev3 Education", "PROP"),
                ("Lego Ev3 Classrooom", "PROP"), ("Gimp", "OS"),
                ("HotPotatoes", "FREE"),
                ("IBM SPSS statistics version 30", "PROP"),
                ("Kompozer", "OS"), ("LibreOffice", "OS"),
                ("Listen N Write", "FREE"),
                ("LightBurn (imprimantes 3D)", "PROP"),
                ("Minicraft education", "PROP"),
                ("MS Office LTSC Pro 2021", "PROP"),
                ("NetLogo", "OS"), ("Netlogo 3D", "OS"),
                ("Notepad++", "OS"), ("NetSupportSchool", "PROP"),
                ("Nvivo 15", "PROP"), ("Nvu", "OS"),
                ("OBS Studio", "OS"), ("OO4Kids", "OS"),
                ("Open-Shot Video editor", "OS"),
                ("OpenStat", "OS"),
                ("Enregistreur Panopto", "PROP"),
                ("Presenter IPEVO", "FREE"), ("Paint 3D", "FREE"),
                ("Qgis", "OS"), ("Quandary", "FREE"),
                ("R", "OS"), ("RStudio", "OS"),
                ("Revo Uninstaller", "FREEM"), ("Shotcut", "OS"),
                ("Smart NoteBook (Lumino)", "PROP"),
                ("Stellarium", "OS"), ("Tinn-R", "OS"),
                ("Twine", "OS"), ("Visualizer LTSE", "FREE"),
                ("VLC", "OS"), ("WEDO Education", "PROP"),
                ("Weft QDA", "OS"), ("Xmind", "FREEM"),
                ("ZOOM", "FREEM"), ("Zotero", "OS"),
                ("Zotero extension", "OS"), ("WinScan2pd", "FREE"),
            ],
        },
    },
    "Sciences de la gestion": {
        "neon": "#ffd700", "icon": "GEST",
        "sections": {
            "PC Windows 11": [
                ("Adobe Acrobat (64-bit)", "PROP"),
                ("Adobe After Effects 2025", "PROP"),
                ("Adobe Creative Cloud", "PROP"),
                ("Adobe Illustrator 2025", "PROP"),
                ("Adobe InDesign 2025", "PROP"),
                ("Adobe Lightroom Classic", "PROP"),
                ("Adobe Media Encoder 2025", "PROP"),
                ("Adobe Photoshop 2025", "PROP"),
                ("Adobe Premiere Pro 2025", "PROP"),
                ("Anaconda3 2024.10-1", "FREEM"),
                ("Miniconda3 py313", "OS"),
                ("Antidote 12", "PROP"),
                ("Apache NetBeans version 26", "OS"),
                ("ArcGIS Pro", "PROP"), ("Arena", "PROP"),
                ("Audacity 3.7.4", "OS"),
                ("Autodesk AutoCAD 2025", "PROP"),
                ("Autodesk Revit 2024", "PROP"),
                ("Beyond 2020 Professional Browser", "PROP"),
                ("BlueJ", "OS"),
                ("CaseWare IDEA 13.0", "PROP"),
                ("EndNote 21", "PROP"), ("Google Chrome", "FREE"),
                ("IBM SPSS Statistics", "PROP"),
                ("IntelliJ IDEA 2025.1.1.1", "FREEM"),
                ("MapInfo Pro 2023 (64-bit)", "PROP"),
                ("MATLAB R2025a", "PROP"),
                ("Microsoft Edge", "FREE"),
                ("Microsoft OneDrive", "FREEM"),
                ("Microsoft PowerBI Desktop", "FREEM"),
                ("Microsoft Project", "PROP"),
                ("Microsoft Visio", "PROP"),
                ("Microsoft Visual Studio Code", "FREE"),
                ("Microsoft Office LTSC Professionnel Plus 2021", "PROP"),
                ("Mozilla Firefox", "OS"),
                ("Notepad++ (64-bit)", "OS"),
                ("NVivo 15", "PROP"),
                ("OpenOffice 4.1.15", "OS"),
                ("PyCharm 2025.1.1.1", "FREEM"),
                ("Python 3.14.2 (64-bit)", "OS"),
                ("QGIS 3.40.7 Bratislava", "OS"),
                ("R for Windows 4.5.0", "OS"),
                ("RStudio", "OS"),
                ("Safe Exam Browser (x64)", "OS"),
                ("SAP GUI for Windows 7.70", "PROP"),
                ("SketchUp 2026", "FREEM"),
                ("Stata 15", "PROP"),
                ("Tableau 2025.2", "PROP"),
                ("Tableau Prep Builder 2025.2", "PROP"),
                ("V-Ray for SketchUp", "PROP"),
                ("Zotero", "OS"), ("Shotcut", "OS"),
            ],
        },
    },
    "Sciences humaines & Langues": {
        "neon": "#ff9f43", "icon": "SHL",
        "sections": {
            "PC Windows 10": [
                ("Adobe Acrobat Reader XI Pro", "PROP"),
                ("Adobe CS3 Illustrator", "PROP"),
                ("Adobe CS3 Photoshop", "PROP"),
                ("Adobe Digital Edition v 4.5.11", "FREE"),
                ("Antidote 12", "PROP"),
                ("ArcGIS Pro v 3.5", "PROP"),
                ("Antconc v 4.2.4", "FREE"),
                ("Audacity v 3.5.1", "OS"),
                ("BlueSky Statistics v 10.3.4", "FREEM"),
                ("Catalyst v 2223", "PROP"),
                ("CloudCompare v 2.13", "OS"),
                ("CMap Tools v 6.04", "FREEM"),
                ("Dictionnaire des difficultes du francais moderne", "PROP"),
                ("Dictionnaire multidictionnaire", "PROP"),
                ("Dictionnaire Le Visuel", "PROP"),
                ("Elan v 6.8", "OS"),
                ("EndNote v 21.5", "PROP"),
                ("Express Scribe v 10.14", "FREEM"),
                ("File Maker Pro 11", "PROP"),
                ("FME Workbench v 2021.0", "PROP"),
                ("Foxit Reader v 2024.3.0.26795", "FREEM"),
                ("Gephi v 0.10", "OS"),
                ("Gimp v 2.10", "OS"),
                ("Google Earth Pro v 7.3.6", "FREE"),
                ("GPower v 3.1.9", "FREE"),
                ("GoldVarb v 3.0b3x", "FREE"),
                ("HyperBase 10", "PROP"),
                ("Inkscape v 1.1", "OS"),
                ("Iramutek v 0.7 alpha 2", "OS"),
                ("Jamovi v 2.2.5", "OS"),
                ("Jasp v 0.13.1", "OS"),
                ("Lexico v 5.13", "PROP"),
                ("Libre Office v 24.8.3.2", "OS"),
                ("Miktex v 4.4", "OS"),
                ("MS Office LTSC Pro 2021", "PROP"),
                ("MyStat v 12", "PROP"),
                ("Navigateur Firefox", "OS"),
                ("Navigateur Google Chrome", "FREE"),
                ("Navigateur Edge", "FREE"),
                ("NetSupportSchool 14", "PROP"),
                ("NotePad ++ v 8.4.4", "OS"),
                ("Nvivo 15", "PROP"),
                ("Pajek XXL et 3XL v 5.17", "FREE"),
                ("PSPP v 1.2.0", "OS"),
                ("Python", "OS"), ("Qgis v 3.28.8", "OS"),
                ("QuickTime 7", "FREE"), ("R v 4.1.1", "OS"),
                ("Revo Uninstaller v 2.4.5", "FREEM"),
                ("RStudio v 2023.06.3 build 561", "OS"),
                ("Sonal v 2.0.99", "OS"), ("SPSS 30", "PROP"),
                ("Spyder v 5.5.0 IDE", "OS"),
                ("Suite NCH v 10.14", "PROP"),
                ("Ucinet v 6.799", "PROP"),
                ("VLC v 3.0.20", "OS"),
                ("Weka v 3.8.3", "OS"),
                ("XMind 8 update 9", "FREEM"),
                ("ZOOM Workplace v 6.2.10", "FREEM"),
                ("Zotero v 6.0.30", "OS"),
                ("Zotero extension dans Firefox", "OS"),
                ("Zotero extension dans Word", "OS"),
                ("Zotero extension dans LibreOffice", "OS"),
            ],
        },
    },
}

# -- Neon colors for labs (used when loading from Excel) ----------------------
LAB_NEON_COLORS = [
    "#00ff88", "#00cfff", "#bf5fff", "#ff6b6b", "#ffd700", "#ff9f43",
    "#e056fd", "#7ed6df", "#f9ca24", "#eb4d4b",
]

LAB_ICONS = [
    "ART", "COM", "SCI", "EDU", "GEST", "SHL",
    "LAB7", "LAB8", "LAB9", "LA10",
]

VALID_LICENCES = {"OS", "PROP", "FREE", "FREEM"}


def load_labs_from_excel(filepath):
    """Charge les donnees des laboratoires depuis un fichier Excel.

    Format attendu du fichier Excel:
    - Chaque feuille (sheet) = un laboratoire
    - Colonnes: Section | Logiciel | Licence
    - La licence doit etre: OS, PROP, FREE ou FREEM
    """
    if not HAS_OPENPYXL:
        raise ImportError(
            "Le module 'openpyxl' est requis.\n"
            "Installez-le avec: pip install openpyxl")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    labs = {}

    for idx, sheet_name in enumerate(wb.sheetnames):
        ws = wb[sheet_name]
        sections = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 3:
                continue
            section = str(row[0]).strip() if row[0] else ""
            logiciel = str(row[1]).strip() if row[1] else ""
            licence = str(row[2]).strip().upper() if row[2] else ""

            if not section or not logiciel:
                continue
            if licence not in VALID_LICENCES:
                licence = "FREE"

            if section not in sections:
                sections[section] = []
            sections[section].append((logiciel, licence))

        if sections:
            neon = LAB_NEON_COLORS[idx % len(LAB_NEON_COLORS)]
            icon = LAB_ICONS[idx % len(LAB_ICONS)] if idx < len(LAB_ICONS) else f"L{idx+1:02d}"
            labs[sheet_name] = {
                "neon": neon,
                "icon": icon,
                "sections": sections,
            }

    wb.close()
    return labs


def load_single_lab_from_excel(filepath):
    """Charge les logiciels d'un seul laboratoire depuis un fichier Excel.

    Format attendu:
    - Colonnes: Section | Logiciel | Licence
    - Lit uniquement la premiere feuille du fichier
    Retourne un dict de sections: {section_name: [(logiciel, licence), ...]}
    """
    if not HAS_OPENPYXL:
        raise ImportError(
            "Le module 'openpyxl' est requis.\n"
            "Installez-le avec: pip install openpyxl")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    sections = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 3:
            continue
        section = str(row[0]).strip() if row[0] else ""
        logiciel = str(row[1]).strip() if row[1] else ""
        licence = str(row[2]).strip().upper() if row[2] else ""

        if not section or not logiciel:
            continue
        if licence not in VALID_LICENCES:
            licence = "FREE"

        if section not in sections:
            sections[section] = []
        sections[section].append((logiciel, licence))

    wb.close()
    return sections


def export_single_lab_to_excel(lab_name, lab_data, filepath):
    """Exporte un seul laboratoire vers un fichier Excel."""
    if not HAS_OPENPYXL:
        raise ImportError(
            "Le module 'openpyxl' est requis.\n"
            "Installez-le avec: pip install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = lab_name[:31]
    ws.append(["Section", "Logiciel", "Licence"])

    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    for sec_name, softs in lab_data["sections"].items():
        for sw_name, sw_lic in softs:
            ws.append([sec_name, sw_name, sw_lic])

    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 10

    wb.save(filepath)
    wb.close()


def export_labs_to_excel(labs, filepath):
    """Exporte les donnees des laboratoires vers un fichier Excel.

    Format du fichier genere:
    - Chaque feuille (sheet) = un laboratoire
    - Colonnes: Section | Logiciel | Licence
    """
    if not HAS_OPENPYXL:
        raise ImportError(
            "Le module 'openpyxl' est requis.\n"
            "Installez-le avec: pip install openpyxl")

    wb = Workbook()
    wb.remove(wb.active)

    for lab_name, lab_data in labs.items():
        ws = wb.create_sheet(title=lab_name[:31])
        ws.append(["Section", "Logiciel", "Licence"])

        # Style header
        for cell in ws[1]:
            cell.font = cell.font.copy(bold=True)

        for sec_name, softs in lab_data["sections"].items():
            for sw_name, sw_lic in softs:
                ws.append([sec_name, sw_name, sw_lic])

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 40
        ws.column_dimensions['C'].width = 10

    wb.save(filepath)
    wb.close()


BG     = "#060610"
PANEL  = "#0c0c1a"
PANEL2 = "#10101e"
BORDER = "#1a1a30"
TEXT   = "#d0dce8"
MUTED  = "#5a6a7a"
TAG_BG = "#14203a"
MATCH  = "#1f3a6e"
WHITE  = "#ffffff"
HEADER_BG = "#050510"
ACCENT = "#00cfff"
ACCENT2 = "#00ff88"
GLOW   = "#00e5ff"


def hex_mix(hex_col, alpha, br=10, bg_=10, bb=15):
    """Melange une couleur hex avec le fond sombre."""
    c = hex_col.lstrip('#')
    r, g, b = int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16)
    return "#{:02x}{:02x}{:02x}".format(
        int(br + (r - br) * alpha),
        int(bg_ + (g - bg_) * alpha),
        int(bb + (b - bb) * alpha))


def neon_btn(parent, text, color, command, active=False):
    """Bouton style neon futuriste avec Frame+Label."""
    bg_col = hex_mix(color, 0.22) if active else PANEL2
    fg_col = color
    bd_col = color if active else hex_mix(color, 0.45)

    outer = tk.Frame(parent, bg=bd_col, padx=1, pady=1)
    inner = tk.Frame(outer, bg=bg_col)
    inner.pack(fill="both", expand=True)

    dot = tk.Label(inner, text="◆", bg=bg_col,
                   fg=color if active else MUTED,
                   font=("Consolas", 8))
    dot.pack(side="left", padx=(6, 2), pady=6)

    lbl = tk.Label(inner, text=text, bg=bg_col,
                   fg=fg_col if active else MUTED,
                   font=("Consolas", 9, "bold"), pady=6, padx=4)
    lbl.pack(side="left", padx=(0, 6))

    def on_enter(_):
        inner.config(bg=hex_mix(color, 0.18))
        dot.config(bg=hex_mix(color, 0.18), fg=color)
        lbl.config(bg=hex_mix(color, 0.18), fg=color)

    def on_leave(_):
        inner.config(bg=bg_col)
        dot.config(bg=bg_col, fg=color if active else MUTED)
        lbl.config(bg=bg_col, fg=fg_col if active else MUTED)

    def on_click(_):
        command()

    for w in (outer, inner, dot, lbl):
        w.bind("<Enter>", on_enter)
        w.bind("<Leave>", on_leave)
        w.bind("<Button-1>", on_click)

    return outer


class LaboSoft(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("LABOSOFT-UQAM // Laboratoires Informatiques")
        self.configure(bg=BG)
        self.geometry("1400x860")
        self.minsize(1000, 650)
        self._labs = dict(LABS)  # copie modifiable
        self._active_lab = None
        self._search_var = tk.StringVar()
        self._filter_licence = None  # None = tous, ou "OS", "PROP", etc.
        self._ready = False
        self._tab_widgets = {}
        self._filter_widgets = {}
        self._build_ui()
        self._ready = True
        self._show_lab(None)
        self._search_var.trace_add("write", self._on_search)
        self._tick()

    # -- Horloge ---------------------------------------------------------------
    def _tick(self):
        self._clock_lbl.config(
            text="⟨ SYS ⟩ " + time.strftime("%Y-%m-%d  //  %H:%M:%S"))
        self.after(1000, self._tick)

    # -- Construction UI -------------------------------------------------------
    def _build_ui(self):
        # ══ HEADER BAR ═══════════════════════════════════════════════════════
        top = tk.Frame(self, bg=HEADER_BG)
        top.pack(fill="x")

        # Glow line top
        tk.Frame(self, bg=ACCENT, height=2).pack(fill="x")

        left = tk.Frame(top, bg=HEADER_BG)
        left.pack(side="left", padx=16, pady=10)

        # Logo icon
        tk.Label(left, text="◈", bg=HEADER_BG, fg=GLOW,
                 font=("Consolas", 22, "bold")).pack(side="left", padx=(0, 10))

        # Title: LABOSOFT-UQAM
        tk.Label(left, text="LABOSOFT", bg=HEADER_BG, fg=WHITE,
                 font=("Consolas", 18, "bold")).pack(side="left")
        tk.Label(left, text="-UQAM", bg=HEADER_BG, fg=GLOW,
                 font=("Consolas", 18, "bold")).pack(side="left")

        # Subtitle
        tk.Label(left, text="  // PRO", bg=HEADER_BG, fg=ACCENT,
                 font=("Consolas", 11, "bold")).pack(side="left", padx=(4, 0))

        right = tk.Frame(top, bg=HEADER_BG)
        right.pack(side="right", padx=16, pady=10)
        self._clock_lbl = tk.Label(right, text="", bg=HEADER_BG, fg=MUTED,
                                   font=("Consolas", 9, "bold"))
        self._clock_lbl.pack(side="left", padx=14)

        # -- Boutons Excel -------------------------------------------------
        if HAS_OPENPYXL:
            excel_frame = tk.Frame(right, bg=HEADER_BG)
            excel_frame.pack(side="left", padx=(0, 12))

            imp_btn = neon_btn(excel_frame, "IMPORTER EXCEL", ACCENT2,
                               command=self._import_excel)
            imp_btn.pack(side="left", padx=3)

            exp_btn = neon_btn(excel_frame, "EXPORTER EXCEL", "#ffd700",
                               command=self._export_excel)
            exp_btn.pack(side="left", padx=3)

        # UQAM badge
        uqam_out = tk.Frame(right, bg=GLOW, padx=1, pady=1)
        uqam_out.pack(side="left")
        uqam_in = tk.Frame(uqam_out, bg=HEADER_BG)
        uqam_in.pack()
        tk.Label(uqam_in, text="◆ UQAM", bg=HEADER_BG, fg=GLOW,
                 font=("Consolas", 9, "bold"), padx=12, pady=5).pack()

        # Glow line under header
        tk.Frame(self, bg=ACCENT, height=1).pack(fill="x")

        # ══ STATUS BAR ═══════════════════════════════════════════════════
        st = tk.Frame(self, bg=PANEL, pady=5, padx=16)
        st.pack(fill="x")
        tk.Label(st, text="◆ SYSTEME OPERATIONNEL", bg=PANEL,
                 fg=ACCENT2, font=("Consolas", 9, "bold")).pack(side="left")
        tk.Label(st, text="   ◆ LABORATOIRES ACTIFS", bg=PANEL,
                 fg=ACCENT, font=("Consolas", 9, "bold")).pack(side="left")
        tk.Label(st, text="   ◆ PRET", bg=PANEL,
                 fg=MUTED, font=("Consolas", 9)).pack(side="left")
        self._count_lbl = tk.Label(st, text="", bg=PANEL, fg=MUTED,
                                   font=("Consolas", 9, "bold"))
        self._count_lbl.pack(side="right")
        tk.Frame(self, bg=BORDER, height=1).pack(fill="x")

        # ══ TAB BAR ══════════════════════════════════════════════════════
        tab_outer = tk.Frame(self, bg=PANEL2, pady=10, padx=12)
        tab_outer.pack(fill="x")

        b = neon_btn(tab_outer, "TOUS LES LABOS", ACCENT,
                     command=lambda: self._show_lab(None), active=True)
        b.pack(side="left", padx=5)
        self._tab_widgets[None] = b

        for name, lab in self._labs.items():
            col = lab["neon"]
            icon = lab["icon"]
            btn = neon_btn(tab_outer, icon, col,
                           command=lambda k=name: self._show_lab(k),
                           active=False)
            btn.pack(side="left", padx=5)
            self._tab_widgets[name] = btn

        tk.Frame(self, bg=BORDER, height=1).pack(fill="x")

        # ══ LICENCE FILTER ════════════════════════════════════════════════
        lic_bar = tk.Frame(self, bg=PANEL2, pady=7, padx=14)
        lic_bar.pack(fill="x")
        tk.Label(lic_bar, text="◈ FILTRE LICENCE:", bg=PANEL2, fg=MUTED,
                 font=("Consolas", 9, "bold")).pack(side="left", padx=(0, 10))

        # Bouton TOUS
        all_btn = neon_btn(lic_bar, "TOUS", ACCENT,
                           command=lambda: self._set_licence_filter(None),
                           active=True)
        all_btn.pack(side="left", padx=4)
        self._filter_widgets[None] = all_btn

        for lic_key, lic_label in LICENCE_LABELS.items():
            col = LICENCE_COLORS[lic_key]
            fb = neon_btn(lic_bar, lic_label, col,
                          command=lambda k=lic_key: self._set_licence_filter(k),
                          active=False)
            fb.pack(side="left", padx=4)
            self._filter_widgets[lic_key] = fb

        # Legende
        legend = tk.Frame(lic_bar, bg=PANEL2)
        legend.pack(side="right")
        for lic_key in ("OS", "PROP", "FREE", "FREEM"):
            col = LICENCE_COLORS[lic_key]
            tk.Label(legend, text="●", bg=PANEL2, fg=col,
                     font=("Consolas", 9)).pack(side="left", padx=(8, 2))
            tk.Label(legend, text=LICENCE_LABELS[lic_key], bg=PANEL2,
                     fg=MUTED, font=("Consolas", 8, "bold")).pack(side="left")

        tk.Frame(self, bg=BORDER, height=1).pack(fill="x")

        # ══ SEARCH BAR ═══════════════════════════════════════════════════
        sb = tk.Frame(self, bg=PANEL2, pady=10, padx=14)
        sb.pack(fill="x")
        tk.Label(sb, text="◈", bg=PANEL2, fg=GLOW,
                 font=("Consolas", 14, "bold")).pack(side="left", padx=(0, 8))

        ef = tk.Frame(sb, bg="#0a0a1a", highlightthickness=2,
                      highlightbackground=ACCENT,
                      highlightcolor=GLOW)
        ef.pack(side="left", fill="x", expand=True, ipady=3)
        self._entry = tk.Entry(ef, textvariable=self._search_var,
                               bg="#0a0a1a", fg=ACCENT,
                               insertbackground=GLOW,
                               relief="flat", font=("Consolas", 13, "bold"))
        self._entry.pack(fill="x", expand=True, padx=10, ipady=5)
        self._entry.insert(0, "Rechercher un logiciel...")
        self._entry.config(fg=MUTED)
        self._entry.bind("<FocusIn>", self._focus_in)
        self._entry.bind("<FocusOut>", self._focus_out)
        tk.Frame(self, bg=BORDER, height=1).pack(fill="x")

        # -- Canvas scroll -----------------------------------------------------
        cont = tk.Frame(self, bg=BG)
        cont.pack(fill="both", expand=True)

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("N.Vertical.TScrollbar",
                        background=BORDER, troughcolor=BG,
                        arrowcolor=ACCENT, bordercolor=BG,
                        darkcolor=BG, lightcolor=BG)

        self._canvas = tk.Canvas(cont, bg=BG, highlightthickness=0)
        vbar = ttk.Scrollbar(cont, orient="vertical",
                             command=self._canvas.yview,
                             style="N.Vertical.TScrollbar")
        self._canvas.configure(yscrollcommand=vbar.set)
        vbar.pack(side="right", fill="y")
        self._canvas.pack(side="left", fill="both", expand=True)

        self._content = tk.Frame(self._canvas, bg=BG)
        self._win_id = self._canvas.create_window(
            (0, 0), window=self._content, anchor="nw")

        self._content.bind("<Configure>",
            lambda e: self._canvas.configure(
                scrollregion=self._canvas.bbox("all")))
        self._canvas.bind("<Configure>",
            lambda e: self._canvas.itemconfig(self._win_id, width=e.width))
        self._canvas.bind_all("<MouseWheel>",
            lambda e: self._canvas.yview_scroll(-1 * (e.delta // 120), "units"))

    # -- Mise a jour onglets ---------------------------------------------------
    def _set_tab_active(self, key):
        for k, outer in self._tab_widgets.items():
            try:
                inner = outer.winfo_children()[0]
                children = inner.winfo_children()
                dot = children[0]
                lbl = children[1]
                col = self._labs[k]["neon"] if k is not None else "#00cfff"
            except Exception:
                continue

            is_active = (k == key)
            bg_col = hex_mix(col, 0.18) if is_active else PANEL2
            bd_col = col if is_active else hex_mix(col, 0.4)
            dot_fg = col if is_active else MUTED
            lbl_fg = col if is_active else MUTED

            outer.config(bg=bd_col)
            inner.config(bg=bg_col)
            dot.config(bg=bg_col, fg=dot_fg)
            lbl.config(bg=bg_col, fg=lbl_fg)

    def _set_filter_btn_active(self, key):
        for k, outer in self._filter_widgets.items():
            try:
                inner = outer.winfo_children()[0]
                children = inner.winfo_children()
                dot = children[0]
                lbl = children[1]
                col = LICENCE_COLORS.get(k, "#00cfff")
            except Exception:
                continue

            is_active = (k == key)
            bg_col = hex_mix(col, 0.18) if is_active else PANEL2
            bd_col = col if is_active else hex_mix(col, 0.4)
            dot_fg = col if is_active else MUTED
            lbl_fg = col if is_active else MUTED

            outer.config(bg=bd_col)
            inner.config(bg=bg_col)
            dot.config(bg=bg_col, fg=dot_fg)
            lbl.config(bg=bg_col, fg=lbl_fg)

    # -- Filtre licence --------------------------------------------------------
    def _set_licence_filter(self, key):
        self._filter_licence = key
        self._set_filter_btn_active(key)
        if not self._ready:
            return
        q = self._search_var.get()
        if q == "Rechercher un logiciel...":
            q = ""
        self._render(q)

    # -- Recherche -------------------------------------------------------------
    def _focus_in(self, _):
        if self._entry.get() == "Rechercher un logiciel...":
            self._entry.delete(0, "end")
            self._entry.config(fg="#00cfff")

    def _focus_out(self, _):
        if self._entry.get().strip() == "":
            self._entry.insert(0, "Rechercher un logiciel...")
            self._entry.config(fg=MUTED)

    def _on_search(self, *_):
        if not self._ready:
            return
        q = self._search_var.get()
        if q == "Rechercher un logiciel...":
            q = ""
        self._render(q)

    # -- Navigation ------------------------------------------------------------
    def _show_lab(self, key):
        self._active_lab = key
        self._set_tab_active(key)
        if not self._ready:
            return
        q = self._search_var.get()
        if q == "Rechercher un logiciel...":
            q = ""
        self._render(q)

    # -- Import / Export Excel -------------------------------------------------
    def _import_excel(self):
        filepath = filedialog.askopenfilename(
            title="Importer un fichier Excel",
            filetypes=[("Fichiers Excel", "*.xlsx *.xls"), ("Tous", "*.*")])
        if not filepath:
            return
        try:
            new_labs = load_labs_from_excel(filepath)
            if not new_labs:
                messagebox.showwarning(
                    "Import Excel",
                    "Aucun laboratoire trouve dans le fichier.\n"
                    "Verifiez le format:\n"
                    "- Chaque feuille = un laboratoire\n"
                    "- Colonnes: Section | Logiciel | Licence")
                return
            self._labs = new_labs
            self._rebuild_tabs()
            self._active_lab = None
            self._show_lab(None)
            count = sum(
                len(s) for lab in new_labs.values()
                for s in lab["sections"].values())
            messagebox.showinfo(
                "Import Excel",
                f"Import reussi!\n"
                f"{len(new_labs)} laboratoire(s), {count} logiciel(s) charges\n"
                f"depuis: {os.path.basename(filepath)}")
        except Exception as e:
            messagebox.showerror("Erreur d'import", str(e))

    def _export_excel(self):
        filepath = filedialog.asksaveasfilename(
            title="Exporter vers Excel",
            defaultextension=".xlsx",
            initialfile="labosoft_data.xlsx",
            filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous", "*.*")])
        if not filepath:
            return
        try:
            export_labs_to_excel(self._labs, filepath)
            messagebox.showinfo(
                "Export Excel",
                f"Export reussi!\nFichier: {os.path.basename(filepath)}\n\n"
                f"Vous pouvez modifier ce fichier dans Excel puis\n"
                f"le re-importer avec le bouton 'IMPORTER EXCEL'.")
        except Exception as e:
            messagebox.showerror("Erreur d'export", str(e))

    def _rebuild_tabs(self):
        """Reconstruit les onglets de laboratoires apres un import."""
        # Supprimer les anciens onglets (sauf "TOUS LES LABOS")
        tous_btn = self._tab_widgets.get(None)
        for key, widget in list(self._tab_widgets.items()):
            if key is not None:
                widget.destroy()
                del self._tab_widgets[key]

        # Trouver le conteneur parent des onglets
        parent = tous_btn.master

        # Recreer les onglets pour chaque lab
        for name, lab in self._labs.items():
            col = lab["neon"]
            icon = lab["icon"]
            btn = neon_btn(parent, icon, col,
                           command=lambda k=name: self._show_lab(k),
                           active=False)
            btn.pack(side="left", padx=4)
            self._tab_widgets[name] = btn

    # -- Per-lab Import / Export -----------------------------------------------
    def _import_lab_excel(self, lab_name):
        """Importer un fichier Excel pour un seul laboratoire."""
        filepath = filedialog.askopenfilename(
            title=f"Importer Excel pour: {lab_name}",
            filetypes=[("Fichiers Excel", "*.xlsx *.xls"), ("Tous", "*.*")])
        if not filepath:
            return
        try:
            sections = load_single_lab_from_excel(filepath)
            if not sections:
                messagebox.showwarning(
                    "Import Excel",
                    "Aucun logiciel trouve dans le fichier.\n"
                    "Verifiez le format:\n"
                    "- Colonnes: Section | Logiciel | Licence")
                return
            self._labs[lab_name]["sections"] = sections
            self._refresh()
            count = sum(len(s) for s in sections.values())
            messagebox.showinfo(
                "Import Excel",
                f"Import reussi!\n"
                f"{count} logiciel(s) charges pour {lab_name}\n"
                f"depuis: {os.path.basename(filepath)}")
        except Exception as e:
            messagebox.showerror("Erreur d'import", str(e))

    def _export_lab_excel(self, lab_name):
        """Exporter un seul laboratoire vers Excel."""
        filepath = filedialog.asksaveasfilename(
            title=f"Exporter Excel pour: {lab_name}",
            defaultextension=".xlsx",
            initialfile=f"{lab_name.replace(' ', '_')}.xlsx",
            filetypes=[("Fichiers Excel", "*.xlsx"), ("Tous", "*.*")])
        if not filepath:
            return
        try:
            export_single_lab_to_excel(
                lab_name, self._labs[lab_name], filepath)
            messagebox.showinfo(
                "Export Excel",
                f"Export reussi!\nFichier: {os.path.basename(filepath)}\n\n"
                f"Modifiez ce fichier dans Excel puis\n"
                f"re-importez-le avec le bouton 'IMPORTER'.")
        except Exception as e:
            messagebox.showerror("Erreur d'export", str(e))

    # -- Add / Delete software -------------------------------------------------
    def _add_software(self, lab_name, section_name=None):
        """Ouvre un dialogue pour ajouter un logiciel."""
        dlg = tk.Toplevel(self)
        dlg.title(f"Ajouter un logiciel - {lab_name}")
        dlg.configure(bg=PANEL)
        dlg.geometry("460x300")
        dlg.resizable(False, False)
        dlg.transient(self)
        dlg.grab_set()

        neon = self._labs[lab_name]["neon"]

        tk.Label(dlg, text="◈ AJOUTER UN LOGICIEL", bg=PANEL, fg=neon,
                 font=("Consolas", 13, "bold")).pack(pady=(14, 10))

        form = tk.Frame(dlg, bg=PANEL)
        form.pack(fill="x", padx=20)

        # Section
        tk.Label(form, text="Section:", bg=PANEL, fg=TEXT,
                 font=("Consolas", 10, "bold"), anchor="w").grid(
                     row=0, column=0, sticky="w", pady=5)
        sec_var = tk.StringVar(value=section_name or "")
        sections = list(self._labs[lab_name]["sections"].keys())
        if sections:
            sec_combo = ttk.Combobox(form, textvariable=sec_var,
                                     values=sections, font=("Consolas", 10),
                                     width=28)
        else:
            sec_combo = tk.Entry(form, textvariable=sec_var,
                                 font=("Consolas", 10), width=30)
        sec_combo.grid(row=0, column=1, sticky="w", pady=5, padx=(10, 0))

        # Logiciel
        tk.Label(form, text="Logiciel:", bg=PANEL, fg=TEXT,
                 font=("Consolas", 10, "bold"), anchor="w").grid(
                     row=1, column=0, sticky="w", pady=5)
        sw_var = tk.StringVar()
        tk.Entry(form, textvariable=sw_var, font=("Consolas", 10),
                 width=30).grid(row=1, column=1, sticky="w", pady=5,
                                padx=(10, 0))

        # Licence
        tk.Label(form, text="Licence:", bg=PANEL, fg=TEXT,
                 font=("Consolas", 10, "bold"), anchor="w").grid(
                     row=2, column=0, sticky="w", pady=5)
        lic_var = tk.StringVar(value="FREE")
        lic_combo = ttk.Combobox(
            form, textvariable=lic_var,
            values=["OS", "PROP", "FREE", "FREEM"],
            font=("Consolas", 10), width=10, state="readonly")
        lic_combo.grid(row=2, column=1, sticky="w", pady=5, padx=(10, 0))

        # Buttons
        btn_frame = tk.Frame(dlg, bg=PANEL)
        btn_frame.pack(pady=16)

        def do_add():
            sec = sec_var.get().strip()
            sw = sw_var.get().strip()
            lic = lic_var.get().strip().upper()
            if not sec or not sw:
                messagebox.showwarning(
                    "Champ manquant",
                    "Veuillez remplir la section et le nom du logiciel.")
                return
            if lic not in VALID_LICENCES:
                lic = "FREE"
            if sec not in self._labs[lab_name]["sections"]:
                self._labs[lab_name]["sections"][sec] = []
            self._labs[lab_name]["sections"][sec].append((sw, lic))
            dlg.destroy()
            self._refresh()

        add_btn = neon_btn(btn_frame, "AJOUTER", neon, command=do_add)
        add_btn.pack(side="left", padx=6)

        cancel_btn = neon_btn(btn_frame, "ANNULER", MUTED,
                              command=dlg.destroy)
        cancel_btn.pack(side="left", padx=6)

    def _delete_software(self, lab_name, sec_name, sw_name, sw_lic):
        """Supprime un logiciel apres confirmation."""
        confirm = messagebox.askyesno(
            "Supprimer un logiciel",
            f"Voulez-vous supprimer '{sw_name}' ({sw_lic})\n"
            f"de la section '{sec_name}' ?")
        if not confirm:
            return
        softs = self._labs[lab_name]["sections"].get(sec_name, [])
        try:
            softs.remove((sw_name, sw_lic))
        except ValueError:
            pass
        # Supprimer la section si elle est vide
        if not softs:
            del self._labs[lab_name]["sections"][sec_name]
        self._refresh()

    def _refresh(self):
        """Rafraichit l'affichage avec les donnees actuelles."""
        q = self._search_var.get()
        if q == "Rechercher un logiciel...":
            q = ""
        self._render(q)

    # -- Rendu -----------------------------------------------------------------
    def _render(self, query=""):
        for w in self._content.winfo_children():
            w.destroy()

        q = query.lower().strip()
        labs = {k: v for k, v in self._labs.items()
                if self._active_lab is None or k == self._active_lab}

        cols = 2 if self._active_lab is None else 1
        wrapper = tk.Frame(self._content, bg=BG)
        wrapper.pack(fill="both", expand=True, padx=14, pady=14)

        # Per-lab action bar (when a single lab is selected)
        if self._active_lab is not None:
            lab_neon = self._labs[self._active_lab]["neon"]

            # Glow separator
            glow_sep = tk.Frame(wrapper, bg=lab_neon, height=2)
            glow_sep.pack(fill="x", pady=(0, 6))

            action_bar = tk.Frame(wrapper, bg=PANEL2, pady=8, padx=12)
            action_bar.pack(fill="x", pady=(0, 12))

            tk.Label(action_bar, text="⟨ GESTION ⟩", bg=PANEL2, fg=lab_neon,
                     font=("Consolas", 10, "bold")).pack(side="left", padx=(0, 12))

            if HAS_OPENPYXL:
                imp_btn = neon_btn(
                    action_bar, "IMPORTER EXCEL", ACCENT2,
                    command=lambda: self._import_lab_excel(self._active_lab))
                imp_btn.pack(side="left", padx=4)

                exp_btn = neon_btn(
                    action_bar, "EXPORTER EXCEL", "#ffd700",
                    command=lambda: self._export_lab_excel(self._active_lab))
                exp_btn.pack(side="left", padx=4)

            add_btn = neon_btn(
                action_bar, "+ AJOUTER LOGICIEL", lab_neon,
                command=lambda: self._add_software(self._active_lab))
            add_btn.pack(side="left", padx=4)

            tk.Frame(wrapper, bg=BORDER, height=1).pack(fill="x", pady=(0, 6))

        col_frames = []
        for c in range(cols):
            f = tk.Frame(wrapper, bg=BG)
            f.pack(side="left", fill="both", expand=True,
                   padx=(0, 10 if c < cols - 1 else 0))
            col_frames.append(f)

        total = 0
        stats = {"OS": 0, "PROP": 0, "FREE": 0, "FREEM": 0}
        for i, (name, lab) in enumerate(labs.items()):
            card_total, card_stats = self._make_card(
                col_frames[i % cols], name, lab, q)
            total += card_total
            for k in stats:
                stats[k] += card_stats.get(k, 0)

        suffix = f'  pour "{query}"' if query else ""
        lic_suffix = ""
        if self._filter_licence:
            lic_suffix = f"  [{LICENCE_LABELS[self._filter_licence]}]"
        stat_str = (f"  (OS:{stats['OS']}  PROP:{stats['PROP']}"
                    f"  FREE:{stats['FREE']}  FREEM:{stats['FREEM']})")
        self._count_lbl.config(
            text=f"◆ {total} logiciel(s) affiche(s){suffix}{lic_suffix}{stat_str}")
        self._canvas.yview_moveto(0)

    def _make_card(self, parent, lab_name, lab, q):
        neon = lab["neon"]
        lic_filter = self._filter_licence
        is_single = self._active_lab is not None

        outer = tk.Frame(parent, bg=neon, padx=1, pady=1)
        outer.pack(fill="x", pady=(0, 12))
        card = tk.Frame(outer, bg=PANEL2)
        card.pack(fill="both", expand=True)

        tk.Frame(card, bg=neon, height=2).pack(fill="x")

        hdr = tk.Frame(card, bg=PANEL2, padx=14, pady=10)
        hdr.pack(fill="x")
        tk.Label(hdr, text="◆", bg=PANEL2, fg=neon,
                 font=("Consolas", 12, "bold")).pack(side="left", padx=(0, 8))
        tk.Label(hdr, text=lab_name.upper(), bg=PANEL2, fg=neon,
                 font=("Consolas", 12, "bold")).pack(side="left")

        tk.Frame(card, bg=BORDER, height=1).pack(fill="x", padx=10)

        body = tk.Frame(card, bg=PANEL2, padx=10, pady=8)
        body.pack(fill="x")

        total = 0
        stats = {"OS": 0, "PROP": 0, "FREE": 0, "FREEM": 0}

        for sec_name, softs in lab["sections"].items():
            filtered = []
            for entry in softs:
                sw_name, sw_lic = entry
                if q and q not in sw_name.lower():
                    continue
                if lic_filter and sw_lic != lic_filter:
                    continue
                filtered.append(entry)

            if not filtered:
                continue

            total += len(filtered)
            for _, lic in filtered:
                stats[lic] = stats.get(lic, 0) + 1

            # Section header with + button
            sr = tk.Frame(body, bg=PANEL2)
            sr.pack(fill="x", pady=(8, 3))
            tk.Label(sr, text=f"  ◈ {sec_name}", bg=PANEL2, fg=neon,
                     font=("Consolas", 10, "bold"),
                     anchor="w").pack(side="left")

            if is_single:
                add_sec_btn = tk.Label(
                    sr, text=" + ", bg=hex_mix(neon, 0.15), fg=neon,
                    font=("Consolas", 10, "bold"), cursor="hand2",
                    padx=5, pady=1)
                add_sec_btn.pack(side="left", padx=8)
                add_sec_btn.bind(
                    "<Button-1>",
                    lambda _, ln=lab_name, sn=sec_name:
                        self._add_software(ln, sn))

            PER_ROW = 3
            for i in range(0, len(filtered), PER_ROW):
                row = tk.Frame(body, bg=PANEL2)
                row.pack(fill="x", pady=1)
                for sw_name, sw_lic in filtered[i:i + PER_ROW]:
                    is_match = bool(q and q in sw_name.lower())
                    lic_col = LICENCE_COLORS.get(sw_lic, TEXT)

                    bg = hex_mix(neon, 0.25) if is_match else TAG_BG
                    fg = neon if is_match else TEXT
                    brd = neon if is_match else BORDER

                    lf = tk.Frame(row, bg=brd, padx=1, pady=1)
                    lf.pack(side="left", padx=2, pady=1)
                    inner_f = tk.Frame(lf, bg=bg)
                    inner_f.pack(fill="both", expand=True)

                    # Nom du logiciel
                    tk.Label(inner_f, text=sw_name, bg=bg, fg=fg,
                             font=("Consolas", 9, "bold"),
                             padx=6, pady=3, width=22,
                             anchor="w").pack(side="left")

                    # Badge licence
                    lic_bg = hex_mix(lic_col, 0.25)
                    tk.Label(inner_f, text=sw_lic, bg=lic_bg, fg=lic_col,
                             font=("Consolas", 7, "bold"),
                             padx=4, pady=2, width=5).pack(side="right",
                                                           padx=(0, 2),
                                                           pady=1)

                    # Delete button (only in single-lab view)
                    if is_single:
                        del_lbl = tk.Label(
                            inner_f, text="✕", bg=bg, fg="#ff4444",
                            font=("Consolas", 9, "bold"), cursor="hand2",
                            padx=3, pady=0)
                        del_lbl.pack(side="right", padx=(0, 2))
                        del_lbl.bind(
                            "<Button-1>",
                            lambda _, ln=lab_name, sn=sec_name,
                                   swn=sw_name, swl=sw_lic:
                                self._delete_software(ln, sn, swn, swl))

        # Badge compteur
        badge_out = tk.Frame(hdr, bg=hex_mix(neon, 0.45), padx=1, pady=1)
        badge_out.pack(side="right")
        tk.Label(badge_out, text=f"  {total} logiciels  ",
                 bg=hex_mix(neon, 0.18), fg=neon,
                 font=("Consolas", 9, "bold"), padx=6, pady=3).pack()

        return total, stats


if __name__ == "__main__":
    print("Demarrage LABOSOFT-UQAM...")
    LaboSoft().mainloop()
